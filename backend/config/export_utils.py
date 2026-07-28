import io
import csv
from django.http import HttpResponse
from config.auth import get_active_company_name


def write_xlsx_header(ws, company_name, title, num_cols):
    """Writes a merged company-name row (bold, 14pt) and title row (bold,
    11pt) above the data, followed by a blank row. Returns the number of
    rows written, so callers building a worksheet manually know the offset
    before appending their own header/data rows."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    last_col = get_column_letter(max(num_cols, 1))
    ws.append([company_name or ''])
    ws.merge_cells(f'A1:{last_col}1')
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([title])
    ws.merge_cells(f'A2:{last_col}2')
    ws['A2'].font = Font(bold=True, size=11)
    ws.append([])
    return 3


def xlsx_response(request, headers, rows, filename, sheet_title='Export'):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    write_xlsx_header(ws, get_active_company_name(request), sheet_title, len(headers))
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response


def pdf_response(request, template_name, context, filename):
    from django.template.loader import render_to_string
    from weasyprint import HTML
    html = render_to_string(template_name, {**context, 'company_name': get_active_company_name(request)})
    pdf = HTML(string=html).write_pdf()
    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}.pdf"'
    return response
