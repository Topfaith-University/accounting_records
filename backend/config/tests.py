import io
from unittest.mock import Mock, patch
from django.test import TestCase
from openpyxl import load_workbook


class GetActiveCompanyNameTests(TestCase):
    def test_returns_company_name_from_validated_jwt(self):
        from config.auth import get_active_company_name
        request = Mock()
        request.auth = {'company_id': 'c1', 'company_name': 'Topfaith University', 'role': 'Admin'}
        self.assertEqual(get_active_company_name(request), 'Topfaith University')

    def test_returns_none_when_request_has_no_auth(self):
        from config.auth import get_active_company_name
        request = Mock()
        request.auth = None
        self.assertIsNone(get_active_company_name(request))


def _request_with_company(company_name):
    request = Mock()
    request.auth = {'company_id': 'c1', 'company_name': company_name, 'role': 'Admin'}
    return request


class XlsxResponseCompanyHeaderTests(TestCase):
    def test_writes_company_name_and_title_above_headers(self):
        from config.export_utils import xlsx_response
        request = _request_with_company('Topfaith University')
        response = xlsx_response(request, ['Code', 'Name'], [['001', 'Cash']], 'accounts', 'Accounts')
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws['A1'].value, 'Topfaith University')
        self.assertEqual(ws['A2'].value, 'Accounts')
        self.assertEqual([c.value for c in ws[4]], ['Code', 'Name'])
        self.assertEqual([c.value for c in ws[5]], ['001', 'Cash'])

    def test_blank_company_name_does_not_error(self):
        from config.export_utils import xlsx_response
        request = _request_with_company(None)
        response = xlsx_response(request, ['Code'], [['001']], 'accounts', 'Accounts')
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        # openpyxl never round-trips a written empty string ('') back as '' —
        # its XML writer emits an empty <c t="inlineStr" /> with no value
        # payload, so load_workbook() always returns None for it. Verified
        # against openpyxl 3.1.5 directly (raw XML inspection). The important
        # behavior under test is that a blank company name does not raise.
        self.assertIsNone(ws['A1'].value)
        self.assertEqual([c.value for c in ws[4]], ['Code'])


class PdfResponseCompanyContextTests(TestCase):
    @patch('weasyprint.HTML')
    @patch('django.template.loader.render_to_string')
    def test_injects_company_name_into_template_context(self, mock_render, mock_html):
        from config.export_utils import pdf_response
        mock_render.return_value = '<html></html>'
        mock_html.return_value.write_pdf.return_value = b'%PDF-1.4'
        request = _request_with_company('Topfaith University')

        pdf_response(request, 'accounts/account_list.html', {'rows': []}, 'accounts')

        context = mock_render.call_args.args[1]
        self.assertEqual(context['company_name'], 'Topfaith University')
        self.assertEqual(context['rows'], [])
