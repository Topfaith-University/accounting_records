import io
from django.http import JsonResponse, HttpResponse
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from config.auth import get_active_company_id
from . import services


def _parse_date(value: str, param_name: str):
    from datetime import date
    try:
        return date.fromisoformat(value)
    except ValueError:
        return None


def _pdf_response(template_name, context, filename):
    from django.template.loader import render_to_string
    from weasyprint import HTML
    html = render_to_string(template_name, context)
    pdf = HTML(string=html).write_pdf()
    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}"'
    return response


def _xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def trial_balance(request):
    company_id = get_active_company_id(request)
    if not company_id:
        return JsonResponse({'detail': 'No active company.'}, status=400)
    date_from = request.query_params.get('date_from', '')
    date_to = request.query_params.get('date_to', '')
    fmt = request.query_params.get('format', 'json')
    if not date_from or not date_to:
        return JsonResponse({'detail': 'date_from and date_to are required.'}, status=400)
    if not _parse_date(date_from, 'date_from') or not _parse_date(date_to, 'date_to'):
        return JsonResponse({'detail': 'Invalid date format. Use YYYY-MM-DD.'}, status=400)

    rows = services.compute_trial_balance(company_id, date_from, date_to)
    ctx = {
        'rows': rows,
        'date_from': date_from,
        'date_to': date_to,
        'total_debits': round(sum(r['total_debits'] for r in rows), 2),
        'total_credits': round(sum(r['total_credits'] for r in rows), 2),
    }

    if fmt == 'pdf':
        pdf_rows = []
        for r in rows:
            d = max(0.0, r['total_debits'] - r['total_credits'])
            c = max(0.0, r['total_credits'] - r['total_debits'])
            pdf_rows.append({**r, 'display_debit': round(d, 2), 'display_credit': round(c, 2)})
        pdf_ctx = {
            'rows': pdf_rows,
            'date_from': date_from,
            'date_to': date_to,
            'total_debits': round(sum(r['display_debit'] for r in pdf_rows), 2),
            'total_credits': round(sum(r['display_credit'] for r in pdf_rows), 2),
        }
        return _pdf_response('reports/trial_balance.html', pdf_ctx, f'trial-balance-{date_from}-{date_to}.pdf')
    if fmt == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = 'Trial Balance'
        ws.append(['Code', 'Account', 'Type', 'Debit (N)', 'Credit (N)'])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            d = max(0.0, r['total_debits'] - r['total_credits'])
            c = max(0.0, r['total_credits'] - r['total_debits'])
            ws.append([r['code'], r['name'], r['account_type'],
                       round(d, 2) if d > 0 else None,
                       round(c, 2) if c > 0 else None])
        total_d = round(sum(max(0.0, r['total_debits'] - r['total_credits']) for r in rows), 2)
        total_c = round(sum(max(0.0, r['total_credits'] - r['total_debits']) for r in rows), 2)
        ws.append(['', 'TOTALS', '', total_d, total_c])
        return _xlsx_response(wb, f'trial-balance-{date_from}-{date_to}.xlsx')
    return JsonResponse(ctx)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def income_statement(request):
    company_id = get_active_company_id(request)
    if not company_id:
        return JsonResponse({'detail': 'No active company.'}, status=400)
    date_from = request.query_params.get('date_from', '')
    date_to = request.query_params.get('date_to', '')
    fmt = request.query_params.get('format', 'json')
    if not date_from or not date_to:
        return JsonResponse({'detail': 'date_from and date_to are required.'}, status=400)
    if not _parse_date(date_from, 'date_from') or not _parse_date(date_to, 'date_to'):
        return JsonResponse({'detail': 'Invalid date format. Use YYYY-MM-DD.'}, status=400)

    data = services.compute_income_statement(company_id, date_from, date_to)

    if fmt == 'pdf':
        return _pdf_response('reports/income_statement.html', data, f'income-statement-{date_from}-{date_to}.pdf')
    if fmt == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = 'Income Statement'
        ws.append(['Account', 'Type', 'Amount (N)'])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.append(['REVENUE', '', ''])
        for r in data['revenue']:
            ws.append([r['name'], r['account_type'], r['balance']])
        ws.append(['Total Revenue', '', data['total_revenue']])
        ws.append(['COST OF SALES', '', ''])
        for r in data['cost_of_sales']:
            ws.append([r['name'], r['account_type'], r['balance']])
        ws.append(['Gross Profit', '', data['gross_profit']])
        ws.append(['EXPENSES', '', ''])
        for r in data['expenses']:
            ws.append([r['name'], r['account_type'], r['balance']])
        ws.append(['Total Expenses', '', data['total_expenses']])
        ws.append(['Net Surplus / (Deficit)', '', data['net_surplus']])
        return _xlsx_response(wb, f'income-statement-{date_from}-{date_to}.xlsx')
    return JsonResponse(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def balance_sheet(request):
    company_id = get_active_company_id(request)
    if not company_id:
        return JsonResponse({'detail': 'No active company.'}, status=400)
    as_of_date = request.query_params.get('as_of_date', '')
    fmt = request.query_params.get('format', 'json')
    if not as_of_date:
        return JsonResponse({'detail': 'as_of_date is required.'}, status=400)
    if not _parse_date(as_of_date, 'as_of_date'):
        return JsonResponse({'detail': 'Invalid date format. Use YYYY-MM-DD.'}, status=400)

    data = services.compute_balance_sheet(company_id, as_of_date)

    if fmt == 'pdf':
        return _pdf_response('reports/balance_sheet.html', data, f'balance-sheet-{as_of_date}.pdf')
    if fmt == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = 'Balance Sheet'
        ws.append(['Account', 'Type', 'Balance (N)'])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for section_label, items, total_key in [
            ('ASSETS', data['assets'], 'total_assets'),
            ('LIABILITIES', data['liabilities'], 'total_liabilities'),
            ('EQUITY', data['equity'], 'total_equity'),
        ]:
            ws.append([section_label, '', ''])
            for r in items:
                ws.append([r['name'], r['account_type'], r['balance']])
            ws.append([f'Total {section_label.title()}', '', data[total_key]])
        return _xlsx_response(wb, f'balance-sheet-{as_of_date}.xlsx')
    return JsonResponse(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def gl_detail(request):
    company_id = get_active_company_id(request)
    if not company_id:
        return JsonResponse({'detail': 'No active company.'}, status=400)
    account_id = request.query_params.get('account_id', '')
    date_from = request.query_params.get('date_from', '')
    date_to = request.query_params.get('date_to', '')
    fmt = request.query_params.get('format', 'json')
    if not account_id or not date_from or not date_to:
        return JsonResponse({'detail': 'account_id, date_from, and date_to are required.'}, status=400)
    if not _parse_date(date_from, 'date_from') or not _parse_date(date_to, 'date_to'):
        return JsonResponse({'detail': 'Invalid date format. Use YYYY-MM-DD.'}, status=400)

    data = services.compute_gl_detail(company_id, account_id, date_from, date_to)
    if data is None:
        return JsonResponse({'detail': 'Account not found.'}, status=404)

    if fmt == 'pdf':
        return _pdf_response(
            'reports/gl_detail.html', data,
            f'gl-detail-{data["account_code"]}-{date_from}-{date_to}.pdf'
        )
    if fmt == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = f'GL {data["account_code"]}'
        ws.append(['Date', 'Reference', 'Description', 'Type', 'Side', 'Amount (N)', 'Running Balance (N)'])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for line in data['lines']:
            ws.append([line['date'], line['reference'], line['entry_description'],
                       line['entry_type'] or 'MANUAL', line['side'], line['amount'], line['running_balance']])
        ws.append(['', '', '', '', 'CLOSING BALANCE', '', data['closing_balance']])
        return _xlsx_response(wb, f'gl-detail-{data["account_code"]}-{date_from}-{date_to}.xlsx')
    return JsonResponse(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard(request):
    from neomodel import db

    company_id = get_active_company_id(request)
    if not company_id:
        return JsonResponse({'detail': 'No active company.'}, status=400)

    # Balances per account type (group by type and normal_balance)
    results, _ = db.cypher_query("""
        MATCH (a:Account {company_id: $company_id}) WHERE a.is_active = true
        OPTIONAL MATCH (e:JournalEntry {company_id: $company_id})-[:HAS_LINE]->(l:JournalLine)-[:AFFECTS_ACCOUNT]->(a)
        WHERE e.status = 'POSTED'
        RETURN a.account_type, a.normal_balance,
               coalesce(sum(CASE WHEN l.side = 'DEBIT' THEN l.amount ELSE 0 END), 0) AS debits,
               coalesce(sum(CASE WHEN l.side = 'CREDIT' THEN l.amount ELSE 0 END), 0) AS credits
    """, {'company_id': company_id})

    ASSET_TYPES = {'Current Assets', 'Non-Current Assets'}
    LIABILITY_TYPES = {'Current Liabilities', 'Non-Current Liabilities'}

    total_assets = 0.0
    total_liabilities = 0.0

    for acc_type, normal_balance, debits, credits in results:
        if normal_balance == 'DEBIT':
            balance = (debits or 0.0) - (credits or 0.0)
        else:
            balance = (credits or 0.0) - (debits or 0.0)
        if acc_type in ASSET_TYPES:
            total_assets += balance
        elif acc_type in LIABILITY_TYPES:
            total_liabilities += balance

    # AP outstanding: POSTED or PARTIAL PurchaseInvoices
    ap_results, _ = db.cypher_query(
        "MATCH (inv:PurchaseInvoice {company_id: $company_id}) WHERE inv.status IN ['POSTED', 'PARTIAL'] "
        "RETURN coalesce(sum(inv.total_amount - inv.amount_paid), 0.0)",
        {'company_id': company_id}
    )
    ap_outstanding = float(ap_results[0][0]) if ap_results else 0.0

    # AR outstanding: POSTED or PARTIAL SalesInvoices
    ar_results, _ = db.cypher_query(
        "MATCH (inv:SalesInvoice {company_id: $company_id}) WHERE inv.status IN ['POSTED', 'PARTIAL'] "
        "RETURN coalesce(sum(inv.total_amount - inv.amount_received), 0.0)",
        {'company_id': company_id}
    )
    ar_outstanding = float(ar_results[0][0]) if ar_results else 0.0

    # Recent 10 POSTED journal entries
    je_results, _ = db.cypher_query("""
        MATCH (e:JournalEntry {company_id: $company_id}) WHERE e.status = 'POSTED'
        RETURN e.entry_id, e.reference, e.date, e.description, e.total_debit
        ORDER BY e.date DESC, e.created_at DESC LIMIT 10
    """, {'company_id': company_id})
    recent_entries = [
        {
            'entry_id': r[0],
            'reference': r[1],
            'date': str(r[2]),
            'description': r[3],
            'total_debit': r[4] or 0.0,
        }
        for r in je_results
    ]

    return JsonResponse({
        'total_assets': round(total_assets, 2),
        'total_liabilities': round(total_liabilities, 2),
        'net_equity': round(total_assets - total_liabilities, 2),
        'ap_outstanding': round(ap_outstanding, 2),
        'ar_outstanding': round(ar_outstanding, 2),
        'recent_entries': recent_entries,
    })
